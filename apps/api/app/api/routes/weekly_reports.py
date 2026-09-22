from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import FileResponse
from pptx import Presentation
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_db, require_project_access
from app.db.models import Change, Decision, Issue, Project, Risk, User, WeeklyReport, WeeklyReportSnapshot
from app.schemas import WeeklyReportCreate, WeeklyReportOut, WeeklyReportUpdate

router = APIRouter(prefix="/projects/{project_id}/weekly-reports", tags=["weekly-reports"])
REPORTS_DIR = Path(__file__).resolve().parents[3] / "generated_reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)


def _format_value(value):
    if value is None:
        return "—"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    return str(value)


def _entity_lines(title: str, items: list) -> list[str]:
    lines = [title]
    if not items:
        lines.append("Отсутствуют данные.")
        return lines

    for index, item in enumerate(items, start=1):
        lines.append(f"{index}. {item.__class__.__name__}")
        for field in item.__table__.columns.keys():
            lines.append(f"- {field}: {_format_value(getattr(item, field))}")
        lines.append("")
    return lines


def _build_pptx_bytes(
    project: Project,
    generated_date: date,
    key_stage_text: str,
    risks: list[Risk],
    issues: list[Issue],
    decisions: list[Decision],
    changes: list[Change],
) -> bytes:
    presentation = Presentation()
    presentation.slide_width = 9144000
    presentation.slide_height = 6858000

    slide_data = [
        (
            "Недельный отчет",
            [
                f"Проект: {project.name}",
                "Разработал: Белоусов Д.А",
                f"Дата формирования: {generated_date.isoformat()}",
                f"Неделя года: {generated_date.isocalendar().week}",
            ],
        ),
        ("Ключевые этапы", [key_stage_text or "Основные этапы проекта за текущую неделю не заполнены."]),
        ("Риски", _entity_lines("Риски", risks)),
        ("Проблемы", _entity_lines("Проблемы", issues)),
        ("Решения", _entity_lines("Решения", decisions)),
        ("Изменения", _entity_lines("Изменения", changes)),
    ]

    for title, body_lines in slide_data:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        for index, line in enumerate(body_lines):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = line
            paragraph.level = 0

    output = Path("/tmp")
    output.mkdir(exist_ok=True)
    archive_path = output / "weekly_report_template.pptx"
    presentation.save(archive_path)
    return archive_path.read_bytes()


@router.get("", response_model=list[WeeklyReportOut])
def list_weekly_reports(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> list[WeeklyReport]:
    require_project_access(project_id, user, db)
    return db.query(WeeklyReport).filter(WeeklyReport.project_id == project_id).all()


@router.post("", response_model=WeeklyReportOut)
def create_weekly_report(
    project_id: int,
    payload: WeeklyReportCreate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> WeeklyReport:
    require_project_access(project_id, user, db)

    weekly_report = WeeklyReport(
        project_id=project_id,
        reporting_date=payload.reporting_date,
        executive_summary=payload.executive_summary,
        key_achievements=payload.key_achievements,
        next_week_activities=payload.next_week_activities,
        management_comments=payload.management_comments,
        status="draft",
        created_by=user.id,
    )
    db.add(weekly_report)
    db.commit()
    db.refresh(weekly_report)
    return weekly_report


@router.post("/generate")
def generate_weekly_report_document(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    require_project_access(project_id, user, db)

    project = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Project not found")

    generated_date = date.today()
    report = db.query(WeeklyReport).filter(WeeklyReport.project_id == project_id, WeeklyReport.reporting_date == generated_date).first()
    if report is None:
        report = WeeklyReport(
            project_id=project_id,
            reporting_date=generated_date,
            executive_summary="Недельный отчет сгенерирован автоматически.",
            key_achievements="Ключевые этапы проекта за неделю.",
            next_week_activities="Планируется продолжение текущих работ.",
            management_comments="Рекомендации руководства не заполнены.",
            status="published",
            created_by=user.id,
        )
        db.add(report)
    else:
        report.executive_summary = "Недельный отчет сгенерирован автоматически."
        report.key_achievements = "Ключевые этапы проекта за неделю."
        report.status = "published"

    risks = (
        db.query(Risk)
        .filter(Risk.project_id == project_id)
        .filter((Risk.impact_level.in_(["high", "critical"])) | (Risk.rating.in_(["high", "critical"])))
        .all()
    )
    issues = (
        db.query(Issue)
        .filter(Issue.project_id == project_id)
        .filter(Issue.severity.in_(["high", "critical"]))
        .all()
    )
    decisions = (
        db.query(Decision)
        .filter(Decision.project_id == project_id)
        .filter(Decision.responsible_person == "Топ менеджмент")
        .all()
    )
    changes = (
        db.query(Change)
        .filter(Change.project_id == project_id)
        .filter(Change.status == "under_review")
        .all()
    )

    key_stage_text = report.key_achievements or "Ключевые этапы проекта за неделю."
    file_name = f"project_{project_id}_{generated_date.isoformat()}.pptx"
    file_path = REPORTS_DIR / file_name
    file_bytes = _build_pptx_bytes(project, generated_date, key_stage_text, risks, issues, decisions, changes)
    file_path.write_bytes(file_bytes)

    db.commit()
    db.refresh(report)
    return {
        "project_id": project_id,
        "reporting_date": generated_date.isoformat(),
        "file_name": file_name,
        "file_path": str(file_path),
        "status": "generated",
    }


@router.get("/generated/{file_name}")
def download_generated_weekly_report(
    project_id: int,
    file_name: str,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> FileResponse:
    require_project_access(project_id, user, db)
    file_path = REPORTS_DIR / file_name
    if not file_path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Weekly report file not found")
    return FileResponse(path=file_path, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", filename=file_name)


@router.get("/{report_id}", response_model=WeeklyReportOut)
def get_weekly_report(
    project_id: int,
    report_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> WeeklyReport:
    require_project_access(project_id, user, db)
    report = db.query(WeeklyReport).filter(WeeklyReport.id == report_id, WeeklyReport.project_id == project_id).first()
    if report is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Weekly report not found")
    return report


@router.patch("/{report_id}", response_model=WeeklyReportOut)
def update_weekly_report(
    project_id: int,
    report_id: int,
    payload: WeeklyReportUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> WeeklyReport:
    require_project_access(project_id, user, db)
    report = db.query(WeeklyReport).filter(WeeklyReport.id == report_id, WeeklyReport.project_id == project_id).first()
    if report is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Weekly report not found")

    for field, value in payload.model_dump(exclude_unset=True).items():
        if value is None:
            continue
        if getattr(report, field) != value:
            setattr(report, field, value)

    db.commit()
    db.refresh(report)
    return report


@router.post("/{report_id}/publish", response_model=WeeklyReportOut)
def publish_weekly_report(
    project_id: int,
    report_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> WeeklyReport:
    require_project_access(project_id, user, db)

    report = db.query(WeeklyReport).filter(WeeklyReport.id == report_id, WeeklyReport.project_id == project_id).first()
    if report is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Weekly report not found")

    report.status = "published"
    report.published_at = datetime.utcnow()

    snapshot = {
        "project_id": project_id,
        "reporting_date": str(report.reporting_date),
        "status": report.status,
        "executive_summary": report.executive_summary,
        "key_achievements": report.key_achievements,
        "next_week_activities": report.next_week_activities,
        "management_comments": report.management_comments,
    }

    db.add(WeeklyReportSnapshot(weekly_report_id=report.id, snapshot_json=str(snapshot)))
    db.commit()
    db.refresh(report)
    return report
